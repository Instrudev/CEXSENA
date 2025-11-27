from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.http import FileResponse
from django.shortcuts import render
import requests
from docx import Document

from .forms import CedulaForm


TARGET_ENTITY = "SENA REGIONAL HUILA GRUPO ADMINISTRATIVO CEFA"


def replace_placeholders(document: Document, data: dict) -> None:
    placeholders = {
        "{{contractor}}": data.get("contractor", ""),
        "{{entity}}": data.get("entity", ""),
        "{{value}}": str(data.get("value", "")),
        "{{object}}": data.get("object", ""),
        "{{process_id}}": data.get("process_id", ""),
        "{{department}}": data.get("department", ""),
        "{{contract_start_date}}": data.get("contract_start_date", ""),
        "{{contract_end_date}}": data.get("contract_end_date", ""),
        "{{url}}": data.get("url", ""),
    }

    def replace_in_paragraph(paragraph) -> None:
        for placeholder, value in placeholders.items():
            if placeholder in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(placeholder, value)

    for paragraph in document.paragraphs:
        replace_in_paragraph(paragraph)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)


def generar_contrato(request):
    form = CedulaForm(request.POST or None)
    context = {"form": form}

    if request.method == "POST" and form.is_valid():
        cedula = form.cleaned_data["cedula"]
        api_url = (
            "https://paco-api-v2-prod.azure-api.net/paco-v2/secop/contract/contractors/"
            f"{cedula}?start_year=2025&end_year=2025&limit=500&sort=value&order=desc"
        )

        try:
            response = requests.get(api_url, timeout=30)
            response.raise_for_status()
            contratos = response.json()
        except (requests.RequestException, ValueError):
            context["error"] = (
                "No se pudo obtener la información desde la API. "
                "Por favor, inténtalo nuevamente más tarde."
            )
            return render(request, "contratos/formulario.html", context)

        contratos_filtrados = [
            contrato for contrato in contratos if contrato.get("entity") == TARGET_ENTITY
        ]

        if not contratos_filtrados:
            context["error"] = (
                "No se encontraron contratos para la cédula ingresada "
                "en la entidad especificada."
            )
            return render(request, "contratos/formulario.html", context)

        primer_contrato = contratos_filtrados[0]
        plantilla_path = Path(settings.BASE_DIR) / "static" / "plantillas" / "plantilla.docx"

        if not plantilla_path.exists():
            context["error"] = "La plantilla de Word no está disponible en la ruta configurada."
            return render(request, "contratos/formulario.html", context)

        document = Document(plantilla_path)
        replace_placeholders(document, primer_contrato)

        output = BytesIO()
        document.save(output)
        output.seek(0)

        filename = f"resultado_{cedula}.docx"
        return FileResponse(output, as_attachment=True, filename=filename)

    return render(request, "contratos/formulario.html", context)
